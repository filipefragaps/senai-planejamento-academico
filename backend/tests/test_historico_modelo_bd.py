"""Relatório Histórico de Aulas no mesmo formato do "Baixar Modelo BD" (e reimportável)."""
import io
import re
from datetime import date, time
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.aula import Aula
from app.models.evento import Evento
from app.models.professor import Professor
from app.services.excel_export import COLUNAS_MODELO_BD, exportar_historico_aulas
from app.services.excel_import_seduc import importar_cronograma_seduc

TEMPLATES_TS = Path(__file__).resolve().parents[2] / "frontend" / "src" / "lib" / "templates.ts"


def _colunas_do_modelo_no_frontend() -> list[str]:
    """Cabeçalho do downloadModeloBancoDados, lido direto do frontend (fonte da verdade do modelo)."""
    ts = TEMPLATES_TS.read_text(encoding="utf-8")
    trecho = ts[ts.index("export function downloadModeloBancoDados"):]
    primeira_linha = re.search(r"aoa_to_sheet\(\[\s*\[(.*?)\]", trecho, re.S).group(1)
    return re.findall(r'"([^"]*)"', primeira_linha)


async def test_cabecalho_identico_ao_modelo_bd():
    assert COLUNAS_MODELO_BD == _colunas_do_modelo_no_frontend()


async def test_historico_exporta_no_formato_do_modelo_e_reimporta(db_session):
    prof = Professor(nome="ALAN JOHNES FELIX DE SOUSA", tipo="Mensalista", horas_contratadas=40, ativo=True)
    db_session.add(prof)
    ev = Evento(nome_turma="926342", disciplina="TÉCNICO EM ELETROTÉCNICA - SEDUC - 21539",
                tipo_modalidade="35 TÉCNICO DE NÍVEL MÉDIO - SEDUC", modalidade="Presencial",
                carga_horaria_total=40, horas_semanais=3, data_inicio=date(2025, 3, 1), data_fim=date(2025, 12, 1),
                dias_semana=[5], horario_inicio=time(19), horario_fim=time(22))
    db_session.add(ev)
    await db_session.flush()
    db_session.add_all([
        Aula(evento_id=ev.id, professor_id=prof.id, data=date(2025, 3, 15), horario_inicio=time(19), horario_fim=time(22),
             turno="Noite", uc_nome_original="Instalação de Sistemas Elétricos Industriais", numero_aula=1,
             ambiente="Lab Elétrica 1", etapa="BÁSICO", status="Agendada"),
        Aula(evento_id=ev.id, data=date(2025, 3, 17), horario_inicio=time(7, 30), horario_fim=time(12),
             turno="Manhã", numero_aula=2, subturma="A", status="Realizada"),
    ])
    await db_session.flush()

    conteudo = await exportar_historico_aulas(db_session, evento_id=ev.id)
    ws = load_workbook(io.BytesIO(conteudo)).active
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    assert ws.title == "Cronograma"
    assert linhas[0] == COLUNAS_MODELO_BD
    assert linhas[1] == ["15/03/2025", "926342", "Noite", "19:00", "22:00",
                         "TÉCNICO EM ELETROTÉCNICA - SEDUC - 21539", "Instalação de Sistemas Elétricos Industriais",
                         1, None, "ALAN JOHNES FELIX DE SOUSA", "Lab Elétrica 1", 3, "BÁSICO",
                         "35 TÉCNICO DE NÍVEL MÉDIO - SEDUC", None, "Agendada"]
    assert linhas[2][3:5] == ["07:30", "12:00"] and linhas[2][11] == 4.5 and linhas[2][8] == "A"

    # Reimportar o próprio histórico atualiza as mesmas aulas (não duplica)
    antes = (await db_session.execute(select(func.count()).select_from(Aula))).scalar_one()
    res = await importar_cronograma_seduc(conteudo, db_session)
    await db_session.flush()
    depois = (await db_session.execute(select(func.count()).select_from(Aula))).scalar_one()
    assert res["erros"] == [] and res["criadas"] == 0 and res["atualizadas"] == 2
    assert depois == antes
