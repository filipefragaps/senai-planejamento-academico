"""
CRM de Gestão de Leads — Meta Ads.
Perfis: admin (visão total + configurações) e vendedor (somente os próprios leads).
"""
import io
import json
import logging
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import JSONResponse, PlainTextResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import and_, case, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import settings
from app.core.deps import get_current_user, require_admin
from app.database import get_db
from app.models.crm import CrmCampanha, CrmHistorico, CrmLead
from app.models.oferta import OfertaCurso
from app.models.usuario import Usuario
from app.services import crm as svc
from app.services import meta_leads as meta

router = APIRouter(prefix="/crm", tags=["CRM de Leads"])
log = logging.getLogger("crm.meta")


async def require_crm(current_user: Usuario = Depends(get_current_user)) -> Usuario:
    if current_user.perfil not in svc.PERFIS_CRM:
        raise HTTPException(403, "Acesso negado: CRM disponível para administrador e vendedor")
    return current_user


# ── Schemas ───────────────────────────────────────────────────────────────────

class LeadCreate(BaseModel):
    nome: str
    telefone: Optional[str] = None
    email: Optional[str] = None
    campanha: Optional[str] = None
    oferta_id: Optional[int] = None
    temperatura: Optional[str] = None
    observacoes: Optional[str] = None
    vendedor_id: Optional[int] = None


class LeadUpdate(BaseModel):
    # vendedor e admin
    temperatura: Optional[str] = None
    status: Optional[str] = None
    observacoes: Optional[str] = None
    ultimo_contato: Optional[date] = None
    proximo_retorno: Optional[date] = None
    motivo_perda: Optional[str] = None
    motivo_perda_comentario: Optional[str] = None
    data_matricula: Optional[date] = None
    vencimento_primeira_parcela: Optional[date] = None
    parcela_paga_em: Optional[date] = None
    oferta_id: Optional[int] = None
    nota: Optional[str] = None
    # somente admin
    vendedor_id: Optional[int] = None
    nome: Optional[str] = None
    telefone: Optional[str] = None
    email: Optional[str] = None
    campanha_id: Optional[int] = None


class AtribuirRequest(BaseModel):
    lead_ids: list[int]
    vendedor_id: Optional[int] = None


class CampanhaCreate(BaseModel):
    nome: str
    oferta_id: Optional[int] = None
    vendedor_padrao_id: Optional[int] = None


class CampanhaUpdate(BaseModel):
    nome: Optional[str] = None
    oferta_id: Optional[int] = None
    vendedor_padrao_id: Optional[int] = None
    ativa: Optional[bool] = None
    aplicar_oferta_leads: bool = False


class ConfigUpdate(BaseModel):
    horas_alerta_novo: Optional[float] = None
    dias_alerta_parcela: Optional[int] = None


CAMPOS_VENDEDOR = {
    "temperatura", "status", "observacoes", "ultimo_contato", "proximo_retorno",
    "motivo_perda", "motivo_perda_comentario", "data_matricula", "oferta_id", "nota",
    "vencimento_primeira_parcela", "parcela_paga_em",
}
# O que o vendedor ainda pode alterar num lead Matriculado (controle da 1ª parcela)
CAMPOS_PARCELA = {"vencimento_primeira_parcela", "parcela_paga_em"}
CAMPOS_ADMIN = CAMPOS_VENDEDOR | {"vendedor_id", "nome", "telefone", "email", "campanha_id"}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _opts():
    return (
        selectinload(CrmLead.campanha),
        selectinload(CrmLead.vendedor),
        selectinload(CrmLead.oferta),
    )


def _rotulo_oferta(o: OfertaCurso | None) -> str | None:
    return f"{o.codigo_evento} – {o.nome_curso}" if o else None


def _serializar(lead: CrmLead, params: svc.ParamsAlerta) -> dict:
    return {
        "id": lead.id,
        "meta_lead_id": lead.meta_lead_id,
        "nome": lead.nome,
        "telefone": lead.telefone,
        "email": lead.email,
        "campanha_id": lead.campanha_id,
        "campanha_nome": lead.campanha.nome if lead.campanha else None,
        "conjunto_anuncio": lead.conjunto_anuncio,
        "anuncio": lead.anuncio,
        "formulario": lead.formulario,
        "plataforma": lead.plataforma,
        "origem": lead.origem,
        "recebido_em": svc.como_utc(lead.recebido_em).isoformat(),
        "dados_extras": lead.dados_extras or {},
        "temperatura": lead.temperatura,
        "status": lead.status,
        "status_label": svc.STATUS_LABEL.get(lead.status, lead.status),
        "finalizado": lead.status in svc.STATUS_FINAIS,
        "vendedor_id": lead.vendedor_id,
        "vendedor_nome": lead.vendedor.nome if lead.vendedor else None,
        "oferta_id": lead.oferta_id,
        "oferta_label": _rotulo_oferta(lead.oferta),
        "observacoes": lead.observacoes,
        "ultimo_contato": lead.ultimo_contato.isoformat() if lead.ultimo_contato else None,
        "proximo_retorno": lead.proximo_retorno.isoformat() if lead.proximo_retorno else None,
        "motivo_perda": lead.motivo_perda,
        "motivo_perda_label": svc.MOTIVO_LABEL.get(lead.motivo_perda) if lead.motivo_perda else None,
        "motivo_perda_comentario": lead.motivo_perda_comentario,
        "data_matricula": lead.data_matricula.isoformat() if lead.data_matricula else None,
        "vencimento_primeira_parcela": lead.vencimento_primeira_parcela.isoformat() if lead.vencimento_primeira_parcela else None,
        "parcela_paga_em": lead.parcela_paga_em.isoformat() if lead.parcela_paga_em else None,
        "parcela_pendente": svc.parcela_pendente(lead),
        "alertas": svc.calcular_alertas(lead, params),
    }


def _periodo_utc(data_inicio: Optional[date], data_fim: Optional[date]):
    """Datas do calendário brasileiro → intervalo UTC [ini, fim)."""
    ini = datetime.combine(data_inicio, time.min, svc.FUSO_BR).astimezone(timezone.utc) if data_inicio else None
    fim = (
        datetime.combine(data_fim + timedelta(days=1), time.min, svc.FUSO_BR).astimezone(timezone.utc)
        if data_fim else None
    )
    return ini, fim


async def _filtrar(
    q, db: AsyncSession, user: Usuario, *,
    busca: Optional[str] = None, campanha_id: Optional[int] = None, status: Optional[str] = None,
    temperatura: Optional[str] = None, vendedor_id: Optional[int] = None, oferta_id: Optional[int] = None,
    data_inicio: Optional[date] = None, data_fim: Optional[date] = None, alerta: Optional[str] = None,
):
    if user.perfil == "vendedor":
        q = q.where(CrmLead.vendedor_id == user.id)
    elif vendedor_id is not None:
        q = q.where(CrmLead.vendedor_id == None if vendedor_id == 0 else CrmLead.vendedor_id == vendedor_id)  # noqa: E711
    if busca:
        b = f"%{busca.strip()}%"
        conds = [CrmLead.nome.ilike(b), CrmLead.email.ilike(b)]
        digitos = svc.normalizar_telefone(busca)
        if digitos:
            conds.append(CrmLead.telefone.like(f"%{digitos}%"))
        q = q.where(or_(*conds))
    if campanha_id:
        q = q.where(CrmLead.campanha_id == campanha_id)
    if oferta_id:
        q = q.where(CrmLead.oferta_id == oferta_id)
    if status:
        q = q.where(CrmLead.status.in_([s for s in status.split(",") if s]))
    if temperatura:
        temps = [t for t in temperatura.split(",") if t]
        conds = [CrmLead.temperatura.in_([t for t in temps if t != "sem"])]
        if "sem" in temps:
            conds.append(CrmLead.temperatura == None)  # noqa: E711
        q = q.where(or_(*conds))
    ini, fim = _periodo_utc(data_inicio, data_fim)
    if ini:
        q = q.where(CrmLead.recebido_em >= ini)
    if fim:
        q = q.where(CrmLead.recebido_em < fim)
    if alerta:
        abertos = CrmLead.status.notin_(svc.STATUS_FINAIS)
        hoje = svc.hoje_br()
        params = await svc.parametros_alerta(db)
        corte = svc.agora() - timedelta(hours=params.horas_novo)
        parcela_pendente = and_(CrmLead.status == "matriculado", CrmLead.parcela_paga_em == None,  # noqa: E711
                                CrmLead.vencimento_primeira_parcela != None)  # noqa: E711
        filtros = {
            "parcela_vencida": and_(parcela_pendente, CrmLead.vencimento_primeira_parcela < hoje),
            "parcela_a_vencer": and_(parcela_pendente, CrmLead.vencimento_primeira_parcela >= hoje,
                                     CrmLead.vencimento_primeira_parcela <= hoje + timedelta(days=params.dias_parcela)),
            "novo_sem_atendimento": and_(CrmLead.status == "novo", CrmLead.recebido_em < corte),
            "retorno_atrasado": and_(abertos, CrmLead.proximo_retorno < hoje),
            "retorno_hoje": and_(abertos, CrmLead.proximo_retorno == hoje),
            "sem_vendedor": and_(abertos, CrmLead.vendedor_id == None),  # noqa: E711
        }
        selecionados = [filtros[a] for a in alerta.split(",") if a in filtros]
        if selecionados:
            q = q.where(or_(*selecionados))
    return q


async def _carregar_lead(db: AsyncSession, lead_id: int, user: Usuario) -> CrmLead:
    lead = (await db.execute(select(CrmLead).options(*_opts()).where(CrmLead.id == lead_id))).scalar_one_or_none()
    if not lead:
        raise HTTPException(404, "Lead não encontrado")
    if user.perfil == "vendedor" and lead.vendedor_id != user.id:
        raise HTTPException(403, "Este lead não está atribuído a você")
    return lead


# ── Opções / configuração ─────────────────────────────────────────────────────

@router.get("/opcoes")
async def opcoes(db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_crm)):
    vendedores = await svc.vendedores_ativos(db)
    return {
        "status": [{"value": v, "label": l, "final": v in svc.STATUS_FINAIS} for v, l in svc.STATUS],
        "temperaturas": [{"value": v, "label": l} for v, l in svc.TEMPERATURAS],
        "motivos_perda": [{"value": v, "label": l} for v, l in svc.MOTIVOS_PERDA],
        "vendedores": [{"id": u.id, "nome": u.nome} for u in vendedores],
        "horas_alerta_novo": await svc.horas_alerta_novo(db),
        "dias_alerta_parcela": await svc.dias_alerta_parcela(db),
    }


@router.get("/config")
async def obter_config(db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    return {"horas_alerta_novo": await svc.horas_alerta_novo(db), "dias_alerta_parcela": await svc.dias_alerta_parcela(db)}


@router.put("/config")
async def salvar_config(data: ConfigUpdate, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    if data.horas_alerta_novo is not None:
        if not (0 < data.horas_alerta_novo <= 720):
            raise HTTPException(400, "Informe um valor entre 0 e 720 horas")
        await svc.set_config(db, "horas_alerta_novo", str(data.horas_alerta_novo))
    if data.dias_alerta_parcela is not None:
        if not (0 <= data.dias_alerta_parcela <= 60):
            raise HTTPException(400, "Informe um valor entre 0 e 60 dias")
        await svc.set_config(db, "dias_alerta_parcela", str(data.dias_alerta_parcela))
    return {"horas_alerta_novo": await svc.horas_alerta_novo(db), "dias_alerta_parcela": await svc.dias_alerta_parcela(db)}


# ── Leads ─────────────────────────────────────────────────────────────────────

@router.get("/leads")
async def listar_leads(
    busca: Optional[str] = None,
    campanha_id: Optional[int] = None,
    status: Optional[str] = Query(default=None, description="lista separada por vírgula"),
    temperatura: Optional[str] = Query(default=None, description="quente,morno,frio,sem"),
    vendedor_id: Optional[int] = Query(default=None, description="0 = sem vendedor"),
    oferta_id: Optional[int] = None,
    data_inicio: Optional[date] = None,
    data_fim: Optional[date] = None,
    alerta: Optional[str] = None,
    ordem: str = Query(default="recentes", description="recentes | antigos | retorno"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
    user: Usuario = Depends(require_crm),
):
    filtros = dict(busca=busca, campanha_id=campanha_id, status=status, temperatura=temperatura,
                   vendedor_id=vendedor_id, oferta_id=oferta_id, data_inicio=data_inicio,
                   data_fim=data_fim, alerta=alerta)
    base = await _filtrar(select(CrmLead), db, user, **filtros)
    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one()

    q = base.options(*_opts())
    if ordem == "antigos":
        q = q.order_by(CrmLead.recebido_em.asc())
    elif ordem == "retorno":
        q = q.order_by(CrmLead.proximo_retorno.is_(None), CrmLead.proximo_retorno.asc(), CrmLead.recebido_em.desc())
    else:
        q = q.order_by(CrmLead.recebido_em.desc())
    q = q.offset((page - 1) * page_size).limit(page_size)
    leads = (await db.execute(q)).scalars().all()

    params = await svc.parametros_alerta(db)
    return {"total": total, "page": page, "page_size": page_size,
            "items": [_serializar(l, params) for l in leads]}


@router.get("/leads/exportar")
async def exportar_leads(
    busca: Optional[str] = None,
    campanha_id: Optional[int] = None,
    status: Optional[str] = None,
    temperatura: Optional[str] = None,
    vendedor_id: Optional[int] = None,
    oferta_id: Optional[int] = None,
    data_inicio: Optional[date] = None,
    data_fim: Optional[date] = None,
    alerta: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: Usuario = Depends(require_crm),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    q = await _filtrar(select(CrmLead), db, user, busca=busca, campanha_id=campanha_id, status=status,
                       temperatura=temperatura, vendedor_id=vendedor_id, oferta_id=oferta_id,
                       data_inicio=data_inicio, data_fim=data_fim, alerta=alerta)
    leads = (await db.execute(q.options(*_opts()).order_by(CrmLead.recebido_em.desc()))).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    cols = [
        ("Recebido em", 18), ("Nome", 30), ("Telefone", 16), ("E-mail", 28), ("Campanha", 30),
        ("Curso (oferta)", 36), ("Vendedor", 22), ("Temperatura", 12), ("Status", 24),
        ("Último contato", 14), ("Próximo retorno", 15), ("Motivo da perda", 28),
        ("Comentário da perda", 30), ("Data matrícula", 14), ("Vencimento 1ª parcela", 16),
        ("1ª parcela paga em", 16), ("Observações", 40),
    ]
    for i, (titulo, larg) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="003B8E")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = larg

    def d(v):
        return v.strftime("%d/%m/%Y") if v else ""

    for r, l in enumerate(leads, 2):
        recebido = svc.como_utc(l.recebido_em).astimezone(svc.FUSO_BR)
        valores = [
            recebido.strftime("%d/%m/%Y %H:%M"), l.nome, l.telefone or "", l.email or "",
            l.campanha.nome if l.campanha else "", _rotulo_oferta(l.oferta) or "",
            l.vendedor.nome if l.vendedor else "", svc.TEMPERATURA_LABEL.get(l.temperatura, ""),
            svc.STATUS_LABEL.get(l.status, l.status), d(l.ultimo_contato), d(l.proximo_retorno),
            svc.MOTIVO_LABEL.get(l.motivo_perda, "") if l.motivo_perda else "",
            l.motivo_perda_comentario or "", d(l.data_matricula), d(l.vencimento_primeira_parcela),
            d(l.parcela_paga_em), l.observacoes or "",
        ]
        for c, v in enumerate(valores, 1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(leads) + 1, 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="leads_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx"'},
    )


@router.get("/leads/{lead_id}")
async def obter_lead(lead_id: int, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_crm)):
    lead = await _carregar_lead(db, lead_id, user)
    hist = (await db.execute(
        select(CrmHistorico).where(CrmHistorico.lead_id == lead_id).order_by(CrmHistorico.id.desc())
    )).scalars().all()
    out = _serializar(lead, await svc.parametros_alerta(db))
    out["historico"] = [
        {
            "id": h.id, "campo": h.campo, "valor_anterior": h.valor_anterior, "valor_novo": h.valor_novo,
            "usuario_nome": h.usuario_nome, "criado_em": svc.como_utc(h.criado_em).isoformat(),
        }
        for h in hist
    ]
    return out


@router.post("/leads", status_code=201)
async def criar_lead_manual(data: LeadCreate, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_crm)):
    if data.temperatura and data.temperatura not in svc.TEMPERATURA_LABEL:
        raise HTTPException(400, "Temperatura inválida")
    dados = data.model_dump()
    if user.perfil == "vendedor":
        dados["vendedor_id"] = user.id
    lead = await svc.criar_lead(db, dados, "manual", user)
    await db.flush()
    lead = await _carregar_lead(db, lead.id, user)
    return _serializar(lead, await svc.parametros_alerta(db))


async def _descrever(db: AsyncSession, campo: str, valor) -> Optional[str]:
    """Valor legível para o histórico."""
    if valor is None or valor == "":
        return None
    if campo == "status":
        return svc.STATUS_LABEL.get(valor, valor)
    if campo == "temperatura":
        return svc.TEMPERATURA_LABEL.get(valor, valor)
    if campo == "motivo_perda":
        return svc.MOTIVO_LABEL.get(valor, valor)
    if campo == "vendedor_id":
        u = await db.get(Usuario, valor)
        return u.nome if u else str(valor)
    if campo == "oferta_id":
        return _rotulo_oferta(await db.get(OfertaCurso, valor)) or str(valor)
    if campo == "campanha_id":
        c = await db.get(CrmCampanha, valor)
        return c.nome if c else str(valor)
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


@router.patch("/leads/{lead_id}")
async def atualizar_lead(
    lead_id: int, data: LeadUpdate, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_crm),
):
    lead = await _carregar_lead(db, lead_id, user)
    enviados = data.model_fields_set
    permitidos = CAMPOS_ADMIN if user.perfil == "admin" else CAMPOS_VENDEDOR
    proibidos = enviados - permitidos
    if proibidos:
        raise HTTPException(403, f"Sem permissão para alterar: {', '.join(sorted(proibidos))}")

    alteracoes = {c: getattr(data, c) for c in enviados if c != "nota"}
    if "telefone" in alteracoes:
        alteracoes["telefone"] = svc.normalizar_telefone(alteracoes["telefone"])
    if "nome" in alteracoes and not (alteracoes["nome"] or "").strip():
        raise HTTPException(400, "O nome não pode ficar vazio")
    alteracoes = {c: v for c, v in alteracoes.items() if getattr(lead, c) != v}

    so_parcela = lead.status == "matriculado" and set(alteracoes) <= CAMPOS_PARCELA
    if lead.status in svc.STATUS_FINAIS and alteracoes and user.perfil != "admin" and not so_parcela:
        raise HTTPException(403, "Lead finalizado: somente o administrador pode corrigir. Você ainda pode registrar anotações"
                                 + (" e o pagamento da 1ª parcela." if lead.status == "matriculado" else "."))

    # Validações de domínio
    if "status" in alteracoes and alteracoes["status"] not in svc.STATUS_LABEL:
        raise HTTPException(400, "Status inválido")
    if alteracoes.get("temperatura") and alteracoes["temperatura"] not in svc.TEMPERATURA_LABEL:
        raise HTTPException(400, "Temperatura inválida")
    if alteracoes.get("motivo_perda") and alteracoes["motivo_perda"] not in svc.MOTIVO_LABEL:
        raise HTTPException(400, "Motivo de perda inválido")
    if alteracoes.get("vendedor_id"):
        v = await db.get(Usuario, alteracoes["vendedor_id"])
        if not v or v.perfil not in svc.PERFIS_CRM or not v.ativo:
            raise HTTPException(400, "Vendedor inválido ou inativo")
    if alteracoes.get("oferta_id") and not await db.get(OfertaCurso, alteracoes["oferta_id"]):
        raise HTTPException(400, "Oferta não encontrada")

    def valor_final(campo):
        return alteracoes[campo] if campo in alteracoes else getattr(lead, campo)

    if alteracoes.get("parcela_paga_em") and valor_final("status") != "matriculado":
        raise HTTPException(400, "O pagamento da 1ª parcela só pode ser registrado para leads matriculados")

    novo_status = valor_final("status")
    if "status" in alteracoes:
        if novo_status == "perdido":
            if not valor_final("motivo_perda"):
                raise HTTPException(400, "Informe o motivo da perda")
            if valor_final("motivo_perda") == "outro" and not (valor_final("motivo_perda_comentario") or "").strip():
                raise HTTPException(400, "Para o motivo 'Outro', descreva o motivo no comentário")
        if novo_status == "matriculado" and not valor_final("data_matricula"):
            raise HTTPException(400, "Informe a data da matrícula")
        if novo_status == "matriculado" and not valor_final("vencimento_primeira_parcela"):
            raise HTTPException(400, "Informe o vencimento da 1ª parcela")
        if novo_status == "aguardando_retorno" and not valor_final("proximo_retorno"):
            raise HTTPException(400, "Informe a data do próximo retorno")
        # Saiu de "Novo lead": marca o primeiro atendimento e o último contato
        if lead.status == "novo" and novo_status != "novo":
            if lead.primeiro_atendimento_em is None:
                lead.primeiro_atendimento_em = svc.agora()
            if "ultimo_contato" not in alteracoes and lead.ultimo_contato is None:
                alteracoes["ultimo_contato"] = svc.hoje_br()

    for campo, valor in alteracoes.items():
        anterior = getattr(lead, campo)
        if campo != "observacoes":
            svc.registrar(db, lead, user, campo, await _descrever(db, campo, anterior), await _descrever(db, campo, valor))
        else:
            svc.registrar(db, lead, user, campo, None, "Observações atualizadas")
        setattr(lead, campo, valor)

    if data.nota and data.nota.strip():
        svc.registrar(db, lead, user, "nota", None, data.nota.strip())

    await db.flush()
    db.expire(lead)
    return await obter_lead(lead_id, db, user)


@router.delete("/leads/{lead_id}", status_code=204)
async def excluir_lead(lead_id: int, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    lead = await db.get(CrmLead, lead_id)
    if not lead:
        raise HTTPException(404, "Lead não encontrado")
    await db.delete(lead)


class ExcluirLeadsRequest(BaseModel):
    lead_ids: list[int]
    senha: str


def _conferir_senha(user: Usuario, senha: Optional[str]) -> None:
    from app.core.security import verify_password
    if not senha or not verify_password(senha, user.hashed_password):
        raise HTTPException(403, "Senha incorreta")


@router.post("/leads/excluir")
async def excluir_leads(data: ExcluirLeadsRequest, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin)):
    """Exclui os leads selecionados (com histórico). Exige a senha do administrador."""
    from sqlalchemy import delete
    if not data.lead_ids:
        raise HTTPException(400, "Nenhum lead selecionado")
    _conferir_senha(user, data.senha)
    await db.execute(delete(CrmHistorico).where(CrmHistorico.lead_id.in_(data.lead_ids)))
    n = (await db.execute(delete(CrmLead).where(CrmLead.id.in_(data.lead_ids)))).rowcount
    log.warning("CRM: %s lead(s) excluído(s) por %s", n, user.email)
    return {"excluidos": n}


@router.post("/leads/atribuir")
async def atribuir_leads(data: AtribuirRequest, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin)):
    if data.vendedor_id:
        v = await db.get(Usuario, data.vendedor_id)
        if not v or v.perfil not in svc.PERFIS_CRM or not v.ativo:
            raise HTTPException(400, "Vendedor inválido ou inativo")
    leads = (await db.execute(select(CrmLead).where(CrmLead.id.in_(data.lead_ids)))).scalars().all()
    novo_nome = await _descrever(db, "vendedor_id", data.vendedor_id)
    alterados = 0
    for lead in leads:
        if lead.vendedor_id == data.vendedor_id:
            continue
        svc.registrar(db, lead, user, "vendedor_id", await _descrever(db, "vendedor_id", lead.vendedor_id), novo_nome)
        lead.vendedor_id = data.vendedor_id
        alterados += 1
    return {"atribuidos": alterados}


@router.post("/leads/distribuir")
async def distribuir_sem_vendedor(db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin)):
    """Distribui leads em aberto sem vendedor (vendedor padrão da campanha ou rodízio)."""
    ativos = await svc.vendedores_ativos(db)
    if not ativos:
        raise HTTPException(400, "Nenhum usuário com perfil vendedor ativo")
    leads = (await db.execute(
        select(CrmLead).options(selectinload(CrmLead.campanha))
        .where(CrmLead.vendedor_id == None, CrmLead.status.notin_(svc.STATUS_FINAIS))  # noqa: E711
        .order_by(CrmLead.recebido_em)
    )).scalars().all()
    nomes = {u.id: u.nome for u in ativos}
    for lead in leads:
        lead.vendedor_id = await svc.escolher_vendedor(db, lead.campanha, ativos)
        svc.registrar(db, lead, user, "vendedor_id", None, nomes.get(lead.vendedor_id))
    return {"distribuidos": len(leads)}


@router.post("/leads/importar-csv")
async def importar_csv(
    arquivo: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: Usuario = Depends(require_admin),
):
    conteudo = await arquivo.read()
    if len(conteudo) > 10 * 1024 * 1024:
        raise HTTPException(400, "Arquivo maior que 10 MB")
    linhas, erros = svc.ler_csv_meta(conteudo)
    if not linhas and erros:
        raise HTTPException(400, erros[0])

    res = await svc.importar_lote(db, linhas, "csv", user)
    return {"inseridos": res["inseridos"], "duplicados": res["duplicados"], "erros": erros[:20], "total_erros": len(erros)}


# ── Painel ────────────────────────────────────────────────────────────────────

def _taxa(parte: int, total: int) -> float:
    return round(parte / total * 100, 1) if total else 0.0


@router.get("/dashboard")
async def dashboard(
    campanha_id: Optional[int] = None,
    status: Optional[str] = None,
    temperatura: Optional[str] = None,
    vendedor_id: Optional[int] = None,
    oferta_id: Optional[int] = None,
    data_inicio: Optional[date] = None,
    data_fim: Optional[date] = None,
    db: AsyncSession = Depends(get_db),
    user: Usuario = Depends(require_crm),
):
    q = await _filtrar(select(CrmLead), db, user, campanha_id=campanha_id, status=status, temperatura=temperatura,
                       vendedor_id=vendedor_id, oferta_id=oferta_id, data_inicio=data_inicio, data_fim=data_fim)
    leads = (await db.execute(q.options(*_opts()))).scalars().all()
    params = await svc.parametros_alerta(db)

    total = len(leads)
    por_status = Counter(l.status for l in leads)
    por_temp = Counter(l.temperatura or "sem" for l in leads)
    alertas = Counter(a for l in leads for a in svc.calcular_alertas(l, params))
    matriculados = por_status.get("matriculado", 0)

    tempos = [
        (svc.como_utc(l.primeiro_atendimento_em) - svc.como_utc(l.recebido_em)).total_seconds() / 3600
        for l in leads if l.primeiro_atendimento_em
    ]

    def agrupar(chave, rotulo):
        grupos: dict = defaultdict(lambda: {"total": 0, "matriculados": 0, "perdidos": 0, "em_aberto": 0})
        nomes: dict = {}
        for l in leads:
            k = chave(l)
            nomes[k] = rotulo(l)
            g = grupos[k]
            g["total"] += 1
            if l.status == "matriculado":
                g["matriculados"] += 1
            elif l.status in svc.STATUS_FINAIS:
                g["perdidos"] += 1
            else:
                g["em_aberto"] += 1
        out = [{"id": k, "nome": nomes[k], **g, "taxa_conversao": _taxa(g["matriculados"], g["total"])}
               for k, g in grupos.items()]
        return sorted(out, key=lambda x: -x["total"])

    por_dia: dict[str, dict] = defaultdict(lambda: {"total": 0, "matriculados": 0})
    for l in leads:
        dia = svc.como_utc(l.recebido_em).astimezone(svc.FUSO_BR).date().isoformat()
        por_dia[dia]["total"] += 1
        if l.status == "matriculado":
            por_dia[dia]["matriculados"] += 1

    motivos = Counter(l.motivo_perda for l in leads if l.status == "perdido" and l.motivo_perda)

    return {
        "total": total,
        "novos": por_status.get("novo", 0),
        "novos_sem_atendimento": alertas.get("novo_sem_atendimento", 0),
        "retornos_atrasados": alertas.get("retorno_atrasado", 0),
        "retornos_hoje": alertas.get("retorno_hoje", 0),
        "sem_vendedor": alertas.get("sem_vendedor", 0),
        "parcelas_a_vencer": alertas.get("parcela_a_vencer", 0),
        "parcelas_vencidas": alertas.get("parcela_vencida", 0),
        "dias_alerta_parcela": params.dias_parcela,
        "matriculas_pagas": sum(1 for l in leads if l.status == "matriculado" and l.parcela_paga_em),
        "matriculas_aguardando_pagamento": sum(1 for l in leads if l.status == "matriculado" and not l.parcela_paga_em),
        "matriculados": matriculados,
        "taxa_conversao": _taxa(matriculados, total),
        "tempo_medio_primeiro_atendimento_h": round(sum(tempos) / len(tempos), 1) if tempos else None,
        "por_status": [
            {"status": s, "label": lbl, "total": por_status.get(s, 0), "final": s in svc.STATUS_FINAIS}
            for s, lbl in svc.STATUS
        ],
        "por_temperatura": [
            {"temperatura": t, "label": lbl, "total": por_temp.get(t, 0)}
            for t, lbl in svc.TEMPERATURAS + [("sem", "Sem classificação")]
        ],
        "por_dia": [{"data": d, **v} for d, v in sorted(por_dia.items())],
        "por_campanha": agrupar(lambda l: l.campanha_id or 0,
                                lambda l: l.campanha.nome if l.campanha else "Sem campanha"),
        "por_curso": agrupar(lambda l: l.oferta_id or 0,
                             lambda l: _rotulo_oferta(l.oferta) or "Sem curso vinculado"),
        "por_vendedor": agrupar(lambda l: l.vendedor_id or 0,
                                lambda l: l.vendedor.nome if l.vendedor else "Sem vendedor"),
        "motivos_perda": [
            {"motivo": m, "label": svc.MOTIVO_LABEL.get(m, m), "total": n} for m, n in motivos.most_common()
        ],
    }


# ── Campanhas ─────────────────────────────────────────────────────────────────

def _serializar_campanha(c: CrmCampanha, total: int = 0, matriculados: int = 0) -> dict:
    return {
        "id": c.id, "nome": c.nome, "meta_campaign_id": c.meta_campaign_id, "ativa": c.ativa,
        "nomes_anteriores": c.nomes_anteriores or [],
        "oferta_id": c.oferta_id, "oferta_label": _rotulo_oferta(c.oferta),
        "vendedor_padrao_id": c.vendedor_padrao_id,
        "vendedor_padrao_nome": c.vendedor_padrao.nome if c.vendedor_padrao else None,
        "total_leads": total, "matriculados": matriculados,
    }


@router.get("/campanhas")
async def listar_campanhas(db: AsyncSession = Depends(get_db), _=Depends(require_crm)):
    camps = (await db.execute(
        select(CrmCampanha).options(selectinload(CrmCampanha.oferta), selectinload(CrmCampanha.vendedor_padrao))
        .order_by(CrmCampanha.ativa.desc(), CrmCampanha.nome)
    )).scalars().all()
    contagem = {
        cid: (tot, mat) for cid, tot, mat in (await db.execute(
            select(CrmLead.campanha_id, func.count(),
                   func.sum(case((CrmLead.status == "matriculado", 1), else_=0)))
            .group_by(CrmLead.campanha_id)
        )).all()
    }
    return [_serializar_campanha(c, *(contagem.get(c.id, (0, 0)))) for c in camps]


async def _validar_campanha(db: AsyncSession, oferta_id, vendedor_id):
    if oferta_id and not await db.get(OfertaCurso, oferta_id):
        raise HTTPException(400, "Oferta não encontrada")
    if vendedor_id:
        v = await db.get(Usuario, vendedor_id)
        if not v or v.perfil not in svc.PERFIS_CRM:
            raise HTTPException(400, "Vendedor inválido")


@router.post("/campanhas", status_code=201)
async def criar_campanha(data: CampanhaCreate, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    nome = data.nome.strip()
    if not nome:
        raise HTTPException(400, "Informe o nome da campanha")
    if (await db.execute(select(CrmCampanha).where(CrmCampanha.nome == nome))).scalar_one_or_none():
        raise HTTPException(400, "Já existe uma campanha com esse nome")
    await _validar_campanha(db, data.oferta_id, data.vendedor_padrao_id)
    c = CrmCampanha(nome=nome, oferta_id=data.oferta_id, vendedor_padrao_id=data.vendedor_padrao_id)
    db.add(c)
    await db.flush()
    await db.refresh(c, ["oferta", "vendedor_padrao"])
    return _serializar_campanha(c)


@router.put("/campanhas/{campanha_id}")
async def atualizar_campanha(
    campanha_id: int, data: CampanhaUpdate, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin),
):
    c = await db.get(CrmCampanha, campanha_id)
    if not c:
        raise HTTPException(404, "Campanha não encontrada")
    enviados = data.model_fields_set
    await _validar_campanha(db, data.oferta_id if "oferta_id" in enviados else None,
                            data.vendedor_padrao_id if "vendedor_padrao_id" in enviados else None)
    if "nome" in enviados:
        novo = (data.nome or "").strip()
        if not novo:
            raise HTTPException(400, "Informe o nome da campanha")
        if novo != c.nome:
            outra = (await db.execute(
                select(CrmCampanha).where(CrmCampanha.nome == novo, CrmCampanha.id != c.id)
            )).scalar_one_or_none()
            if outra:
                raise HTTPException(400, f"Já existe outra campanha chamada “{novo}”")
            # Guarda o nome antigo para reconhecer importações futuras (CSV) que ainda o usam
            anteriores = [n for n in (c.nomes_anteriores or []) if n.lower() != novo.lower()]
            if c.nome.lower() not in [n.lower() for n in anteriores]:
                anteriores.append(c.nome)
            c.nomes_anteriores = anteriores
            c.nome = novo
    for campo in ("oferta_id", "vendedor_padrao_id", "ativa"):
        if campo in enviados:
            setattr(c, campo, getattr(data, campo))

    aplicados = 0
    if data.aplicar_oferta_leads and c.oferta_id:
        leads = (await db.execute(
            select(CrmLead).where(CrmLead.campanha_id == c.id, CrmLead.oferta_id == None)  # noqa: E711
        )).scalars().all()
        rotulo = await _descrever(db, "oferta_id", c.oferta_id)
        for lead in leads:
            lead.oferta_id = c.oferta_id
            svc.registrar(db, lead, user, "oferta_id", None, rotulo)
        aplicados = len(leads)

    await db.flush()
    await db.refresh(c, ["oferta", "vendedor_padrao"])
    return {**_serializar_campanha(c), "leads_atualizados": aplicados}


# ── Integração Meta (Lead Ads) ────────────────────────────────────────────────

@router.get("/webhook/meta", include_in_schema=False)
async def webhook_meta_verificar(
    modo: Optional[str] = Query(default=None, alias="hub.mode"),
    token: Optional[str] = Query(default=None, alias="hub.verify_token"),
    desafio: Optional[str] = Query(default=None, alias="hub.challenge"),
):
    """Verificação feita pela Meta ao cadastrar o webhook: devolve o hub.challenge."""
    if modo == "subscribe" and settings.META_VERIFY_TOKEN and token == settings.META_VERIFY_TOKEN:
        return PlainTextResponse(desafio or "")
    raise HTTPException(403, "Token de verificação inválido")


@router.post("/webhook/meta", include_in_schema=False)
async def webhook_meta_receber(request: Request, db: AsyncSession = Depends(get_db)):
    """Aviso de lead novo. Público, protegido pela assinatura X-Hub-Signature-256."""
    corpo = await request.body()
    if not meta.assinatura_valida(corpo, request.headers.get("X-Hub-Signature-256")):
        raise HTTPException(403, "Assinatura inválida")
    try:
        payload = json.loads(corpo)
    except ValueError:
        raise HTTPException(400, "JSON inválido")
    try:
        res = await meta.processar_webhook(db, payload)
    except (meta.MetaErro, IntegrityError) as e:
        # Responde 200 para a Meta não desativar o webhook; a sincronização periódica recupera o lead
        await db.rollback()
        msg = str(e) if isinstance(e, meta.MetaErro) else "Lead já recebido por outra rotina"
        await meta.registrar_erro(db, f"Webhook: {msg}")
        log.warning("Webhook Meta: %s", msg)
        return {"ok": False}
    return {"ok": True, **res}


def _url_webhook(request: Request) -> str:
    """Endereço público do webhook: PUBLIC_API_URL ou o host desta requisição (atrás do proxy da Railway)."""
    proto = request.headers.get("x-forwarded-proto", request.url.scheme)
    base = (settings.PUBLIC_API_URL or f"{proto}://{request.headers.get('host', request.url.netloc)}").rstrip("/")
    return f"{base}{settings.API_V1_STR}/crm/webhook/meta"


@router.get("/meta/status")
async def meta_status(request: Request, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    return {
        "configurado": meta.configurado(),
        "variaveis": meta.situacao_variaveis(),
        "versao_graph": settings.META_GRAPH_VERSION,
        "sincronizacao_minutos": settings.META_SYNC_MINUTOS,
        "webhook_caminho": f"{settings.API_V1_STR}/crm/webhook/meta",
        "webhook_url": _url_webhook(request),
        **(await meta.estado(db)),
    }


def _erro_meta(e: Exception) -> JSONResponse:
    return JSONResponse(status_code=502, content={"detail": str(e)})


@router.post("/meta/testar")
async def meta_testar(request: Request, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    """Confere token, página, formulários e se a página está assinada para enviar leads ao app."""
    if not meta.configurado():
        raise HTTPException(400, "Defina META_ACCESS_TOKEN e META_PAGE_ID nas variáveis do servidor")
    try:
        async with meta.MetaCliente() as cli:
            try:
                pagina = await cli.conectar_pagina()
                forms = await cli.formularios()
                apps = await cli.assinaturas()
                webhook_app = None
                try:  # informativo: não impede o teste
                    app = await cli.app()
                    webhook_app = {"app": app, **((await cli.webhook_do_app(app["id"])) or {"url": None, "leadgen": False})}
                    webhook_app["url_correta"] = webhook_app.get("url") == _url_webhook(request)
                except meta.MetaErro as e:
                    webhook_app = {"erro": str(e)}
            except meta.MetaErro as e:
                diag = await cli.diagnostico()
                await meta.registrar_erro(db, f"Teste de conexão: {e}")
                return JSONResponse(status_code=502, content={"detail": str(e), "diagnostico": diag})
    except meta.MetaErro as e:
        await meta.registrar_erro(db, f"Teste de conexão: {e}")
        return _erro_meta(e)
    assinada = any("leadgen" in (a.get("subscribed_fields") or []) for a in apps)
    return {
        "pagina": pagina,
        "formularios": [{"id": f["id"], "nome": f.get("name"), "status": f.get("status")} for f in forms],
        "webhook_assinado": assinada,
        "webhook_app": webhook_app,
    }


@router.post("/meta/webhook-app")
async def meta_cadastrar_webhook_app(request: Request, db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    """Cadastra o webhook no app da Meta (objeto Página, campo leadgen) sem usar o painel da Meta."""
    if not meta.configurado() or not settings.META_APP_SECRET or not settings.META_VERIFY_TOKEN:
        raise HTTPException(400, "Defina META_ACCESS_TOKEN, META_PAGE_ID, META_APP_SECRET e META_VERIFY_TOKEN")
    url = _url_webhook(request)
    if not url.startswith("https://"):
        raise HTTPException(400, f"A Meta só aceita endereço https; este servidor está em {url}")
    try:
        async with meta.MetaCliente() as cli:
            app = await cli.app()
            ok = await cli.cadastrar_webhook_do_app(app["id"], url)
            situacao = await cli.webhook_do_app(app["id"])
    except meta.MetaErro as e:
        await meta.registrar_erro(db, f"Cadastro do webhook: {e}")
        return _erro_meta(e)
    return {"cadastrado": ok, "app": app, "url": url, "webhook": situacao}


@router.post("/meta/assinar")
async def meta_assinar(db: AsyncSession = Depends(get_db), _=Depends(require_admin)):
    """Assina a página para enviar os avisos de lead (campo leadgen) ao app."""
    if not meta.configurado():
        raise HTTPException(400, "Defina META_ACCESS_TOKEN e META_PAGE_ID nas variáveis do servidor")
    try:
        async with meta.MetaCliente() as cli:
            await cli.conectar_pagina()
            ok = await cli.assinar_leadgen()
    except meta.MetaErro as e:
        await meta.registrar_erro(db, f"Assinatura da página: {e}")
        return _erro_meta(e)
    return {"assinado": ok}


@router.post("/meta/sincronizar")
async def meta_sincronizar(
    dias: Optional[int] = Query(default=None, ge=1, le=90, description="Buscar os últimos N dias (padrão: desde a última sincronização)"),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_admin),
):
    if not meta.configurado():
        raise HTTPException(400, "Defina META_ACCESS_TOKEN e META_PAGE_ID nas variáveis do servidor")
    try:
        return await meta.sincronizar(db, dias)
    except meta.MetaErro as e:
        await db.rollback()
        await meta.registrar_erro(db, f"Sincronização: {e}")
        return _erro_meta(e)


# ── Limpeza do banco do CRM ───────────────────────────────────────────────────

class LimparCrmRequest(BaseModel):
    senha: str
    confirmacao: str
    incluir_campanhas: bool = False


@router.post("/limpar")
async def limpar_crm(data: LimparCrmRequest, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin)):
    """Apaga todos os leads e o histórico (e, opcionalmente, as campanhas). Exige a senha do administrador."""
    from sqlalchemy import delete
    from app.core.security import verify_password

    if data.confirmacao.strip().upper() != "EXCLUIR":
        raise HTTPException(400, "Digite EXCLUIR para confirmar")
    if not verify_password(data.senha, user.hashed_password):
        raise HTTPException(403, "Senha incorreta")

    historico = (await db.execute(delete(CrmHistorico))).rowcount
    leads = (await db.execute(delete(CrmLead))).rowcount
    campanhas = (await db.execute(delete(CrmCampanha))).rowcount if data.incluir_campanhas else 0
    await svc.set_config(db, "rodizio_ultimo_id", None)
    log.warning("CRM limpo por %s: %s lead(s), %s campanha(s)", user.email, leads, campanhas)
    return {"leads": leads, "historico": historico, "campanhas": campanhas}


class ExcluirCampanhaRequest(BaseModel):
    excluir_leads: bool = False
    senha: Optional[str] = None


@router.post("/campanhas/{campanha_id}/excluir")
async def excluir_campanha(
    campanha_id: int, data: ExcluirCampanhaRequest, db: AsyncSession = Depends(get_db), user: Usuario = Depends(require_admin),
):
    """Exclui a campanha. Os leads dela são excluídos (exige senha) ou ficam sem campanha."""
    from sqlalchemy import delete, update
    c = await db.get(CrmCampanha, campanha_id)
    if not c:
        raise HTTPException(404, "Campanha não encontrada")
    ids = list((await db.execute(select(CrmLead.id).where(CrmLead.campanha_id == campanha_id))).scalars().all())
    if data.excluir_leads and ids:
        _conferir_senha(user, data.senha)
        await db.execute(delete(CrmHistorico).where(CrmHistorico.lead_id.in_(ids)))
        await db.execute(delete(CrmLead).where(CrmLead.id.in_(ids)))
        excluidos, desvinculados = len(ids), 0
    else:
        for lead_id in ids:
            db.add(CrmHistorico(lead_id=lead_id, usuario_id=user.id, usuario_nome=user.nome, campo="campanha_id",
                                valor_anterior=c.nome, valor_novo=None, criado_em=svc.agora()))
        await db.execute(update(CrmLead).where(CrmLead.campanha_id == campanha_id).values(campanha_id=None))
        excluidos, desvinculados = 0, len(ids)
    await db.delete(c)
    log.warning("CRM: campanha '%s' excluída por %s (%s lead(s) excluído(s), %s sem campanha)",
                c.nome, user.email, excluidos, desvinculados)
    return {"campanha": c.nome, "leads_excluidos": excluidos, "leads_sem_campanha": desvinculados}
